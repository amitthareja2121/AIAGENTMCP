from openpyxl import load_workbook
from pathlib import Path
p = Path(r'C:/Users/1000528/files_claude/E2E_Test_Execution_Summary.xlsx')
print('exists=', p.exists())
if p.exists():
    wb = load_workbook(p)
    ws = wb.active
    print(list(ws.iter_rows(values_only=True)))
