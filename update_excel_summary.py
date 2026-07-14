import argparse
from pathlib import Path
from datetime import date, datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import DoughnutChart, BarChart, Reference, Series

BASE_DIR = Path(r'C:\Users\1000528\files_claude')
TARGET_URL = 'https://practicesoftwaretesting.com'

C_HEADER_BG  = '1F3864'
C_HEADER_FG  = 'FFFFFF'
C_TITLE_BG   = '2E75B6'
C_TITLE_FG   = 'FFFFFF'
C_SUMMARY_BG = 'D6E4F0'
C_PASS_BG    = 'E2EFDA'
C_FAIL_BG    = 'FCE4D6'
C_PASS_FG    = '375623'
C_FAIL_FG    = '9C0006'
C_ALT_ROW    = 'F5F5F5'

HEADERS = ['Suite', 'Test Case', 'Test Name', 'Target URL', 'Status', 'Execution Date', 'Notes']
NCOLS   = len(HEADERS)   # 7

def thin_border():
    s = Side(style='thin', color='BFBFBF')
    return Border(left=s, right=s, top=s, bottom=s)

def make_fill(hex_color):
    return PatternFill('solid', fgColor=hex_color)

def last_col_letter():
    return get_column_letter(NCOLS)

def style_title_row(ws, row):
    ws.merge_cells(f'A{row}:{last_col_letter()}{row}')
    cell = ws.cell(row=row, column=1,
                   value=f'E2E Test Execution Report  \u00b7  All Suites  \u00b7  {datetime.today().strftime("%Y-%m-%d")}')
    cell.font      = Font(bold=True, size=14, color=C_TITLE_FG, name='Calibri')
    cell.fill      = make_fill(C_TITLE_BG)
    cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[row].height = 28

def style_summary_row(ws, row, total, passed, failed):
    rate = f'{int(passed/total*100)}%' if total else '\u2014'
    ws.merge_cells(f'A{row}:{last_col_letter()}{row}')
    cell = ws.cell(row=row, column=1,
                   value=f'Total: {total}     \u2705 Passed: {passed}     \u274c Failed: {failed}     Pass Rate: {rate}')
    cell.font      = Font(bold=True, size=11, color='1F3864', name='Calibri')
    cell.fill      = make_fill(C_SUMMARY_BG)
    cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[row].height = 22

def style_header_row(ws, row):
    for col, h in enumerate(HEADERS, start=1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font      = Font(bold=True, size=11, color=C_HEADER_FG, name='Calibri')
        cell.fill      = make_fill(C_HEADER_BG)
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border    = thin_border()
    ws.row_dimensions[row].height = 20

def style_data_row(ws, row, values, is_alt):
    status = values[4]   # Status is 5th column (index 4)
    if status in ('Passed', 'Failed'):
        row_fill = make_fill(C_PASS_BG if status == 'Passed' else C_FAIL_BG)
    elif is_alt:
        row_fill = make_fill(C_ALT_ROW)
    else:
        row_fill = None
    for col, v in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col, value=v)
        cell.font      = Font(size=10, name='Calibri')
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        cell.border    = thin_border()
        if row_fill:
            cell.fill = row_fill
        if col == 5 and status in ('Passed', 'Failed'):
            cell.font = Font(bold=True, size=10, name='Calibri',
                             color=C_PASS_FG if status == 'Passed' else C_FAIL_FG)
    ws.row_dimensions[row].height = 40

def set_column_widths(ws):
    widths = {'A': 10, 'B': 10, 'C': 28, 'D': 38, 'E': 10, 'F': 16, 'G': 55}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

def add_dashboard(wb, all_records):
    total  = len(all_records)
    passed = sum(1 for r in all_records if r['Status'] == 'Passed')
    failed = total - passed

    if 'Dashboard' in wb.sheetnames:
        del wb['Dashboard']
    ds = wb.create_sheet('Dashboard', 0)
    ds.sheet_view.showGridLines = False

    # Title
    ds.merge_cells('A1:H1')
    t = ds.cell(1, 1, 'Test Execution Dashboard')
    t.font      = Font(bold=True, size=16, color=C_TITLE_FG, name='Calibri')
    t.fill      = make_fill(C_TITLE_BG)
    t.alignment = Alignment(horizontal='center', vertical='center')
    ds.row_dimensions[1].height = 32

    # Overall KPIs
    kpis = [('Total Runs', total, 'BDD7EE'), ('Passed', passed, 'E2EFDA'), ('Failed', failed, 'FCE4D6')]
    ds.row_dimensions[3].height = 14
    for col, (label, val, bg) in enumerate(kpis, start=1):
        lc = ds.cell(4, col, label)
        lc.font = Font(bold=True, size=10, color='1F3864', name='Calibri')
        lc.fill = make_fill(bg); lc.alignment = Alignment(horizontal='center')
        lc.border = thin_border(); ds.row_dimensions[4].height = 18
        vc = ds.cell(5, col, val)
        vc.font = Font(bold=True, size=20, color='1F3864', name='Calibri')
        vc.fill = make_fill(bg); vc.alignment = Alignment(horizontal='center', vertical='center')
        vc.border = thin_border(); ds.row_dimensions[5].height = 36
        ds.column_dimensions[get_column_letter(col)].width = 18

    rate_val = f'{int(passed/total*100)}%' if total else '0%'
    ds.merge_cells('D4:F4')
    r4 = ds.cell(4, 4, 'Pass Rate')
    r4.font = Font(bold=True, size=10, color='1F3864', name='Calibri')
    r4.fill = make_fill('FFF2CC'); r4.alignment = Alignment(horizontal='center'); r4.border = thin_border()
    ds.merge_cells('D5:F5')
    r5 = ds.cell(5, 4, rate_val)
    r5.font = Font(bold=True, size=20, color='375623' if failed == 0 else '9C0006', name='Calibri')
    r5.fill = make_fill('FFF2CC'); r5.alignment = Alignment(horizontal='center', vertical='center')
    r5.border = thin_border()
    for c in ['D','E','F']:
        ds.column_dimensions[c].width = 10

    # Donut chart label
    ds.merge_cells('A6:F6')
    cl = ds.cell(6, 1, 'Pass / Fail Distribution')
    cl.font = Font(bold=True, size=11, color='1F3864', name='Calibri')
    cl.fill = make_fill('BDD7EE'); cl.alignment = Alignment(horizontal='center', vertical='center')
    ds.row_dimensions[6].height = 20

    # Chart data in cols H-I — text made white so invisible, chart still reads it
    for r, val in [(3,'Status'),(4,'Passed'),(5,'Failed')]:
        c = ds.cell(r, 8, val)
        c.font = Font(color='FFFFFF', size=1, name='Calibri')
    for r, val in [(3,'Count'),(4,passed),(5,failed)]:
        c = ds.cell(r, 9, val)
        c.font = Font(color='FFFFFF', size=1, name='Calibri')
    ds.column_dimensions['H'].width = 8
    ds.column_dimensions['I'].width = 8

    chart = DoughnutChart()
    chart.style = 10; chart.holeSize = 50; chart.width = 14; chart.height = 10
    data   = Reference(ds, min_col=9, min_row=3, max_row=5)
    labels = Reference(ds, min_col=8, min_row=4, max_row=5)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(labels)
    ds.add_chart(chart, 'A7')

    # ── Suite Breakdown Table ──────────────────────────────────────────────────
    suites = {}
    for r in all_records:
        s = r.get('Suite', 'Unknown')
        if s not in suites:
            suites[s] = {'total': 0, 'passed': 0, 'failed': 0}
        suites[s]['total'] += 1
        if r['Status'] == 'Passed':
            suites[s]['passed'] += 1
        else:
            suites[s]['failed'] += 1

    start_row = 30
    ds.merge_cells(f'A{start_row}:F{start_row}')
    sh = ds.cell(start_row, 1, 'Suite Breakdown')
    sh.font = Font(bold=True, size=12, color=C_TITLE_FG, name='Calibri')
    sh.fill = make_fill('1F3864'); sh.alignment = Alignment(horizontal='center', vertical='center')
    ds.row_dimensions[start_row].height = 22

    hdr_row = start_row + 1
    for col, hdr in enumerate(['Suite', 'Total', 'Passed', 'Failed', 'Pass Rate'], start=1):
        c = ds.cell(hdr_row, col, hdr)
        c.font = Font(bold=True, size=10, color=C_HEADER_FG, name='Calibri')
        c.fill = make_fill(C_HEADER_BG)
        c.alignment = Alignment(horizontal='center'); c.border = thin_border()
    ds.row_dimensions[hdr_row].height = 18

    for i, (suite, stats) in enumerate(sorted(suites.items())):
        drow = hdr_row + 1 + i
        rate = f"{int(stats['passed']/stats['total']*100)}%" if stats['total'] else '0%'
        vals = [suite, stats['total'], stats['passed'], stats['failed'], rate]
        is_pass = stats['failed'] == 0
        for col, v in enumerate(vals, start=1):
            c = ds.cell(drow, col, v)
            c.font = Font(size=10, name='Calibri')
            c.alignment = Alignment(horizontal='center', vertical='center')
            c.border = thin_border()
            if col == 5:
                c.fill = make_fill(C_PASS_BG if is_pass else C_FAIL_BG)
                c.font = Font(bold=True, size=10, name='Calibri',
                              color=C_PASS_FG if is_pass else C_FAIL_FG)
        ds.row_dimensions[drow].height = 20
    for c in ['A','B','C','D','E']:
        ds.column_dimensions[c].width = 16

def rebuild_workbook(path, all_records):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Execution Summary'
    ws.sheet_view.showGridLines = False

    total  = len(all_records)
    passed = sum(1 for r in all_records if r['Status'] == 'Passed')
    failed = total - passed

    style_title_row(ws, 1)
    style_summary_row(ws, 2, total, passed, failed)
    ws.row_dimensions[3].height = 6
    style_header_row(ws, 4)

    for i, rec in enumerate(all_records):
        values = [rec.get(h, '') for h in HEADERS]
        style_data_row(ws, 5 + i, values, is_alt=(i % 2 == 1))

    set_column_widths(ws)
    ws.freeze_panes = 'A5'
    ws.auto_filter.ref = f'A4:{last_col_letter()}4'

    add_dashboard(wb, all_records)
    wb.active = wb['Execution Summary']
    wb.save(path)

def load_existing_records(path):
    wb = load_workbook(path)
    ws = wb['Execution Summary'] if 'Execution Summary' in wb.sheetnames else wb.active
    records = []
    for row in ws.iter_rows(min_row=5, values_only=True):
        if any(v for v in row):
            records.append({HEADERS[i]: (row[i] or '') for i in range(min(len(HEADERS), len(row)))})
    return records

def parse_args():
    parser = argparse.ArgumentParser()
    parser.add_argument('--id',     required=True,  help='TC number e.g. 001')
    parser.add_argument('--suite',  required=False, default='', help='Suite e.g. TS-001')
    parser.add_argument('--status', required=True,  choices=['Passed', 'Failed'])
    parser.add_argument('--scenario', required=False, default='')
    parser.add_argument('--notes',  required=False, default='')
    return parser.parse_args()

def build_record(args):
    id_padded  = str(args.id).zfill(3)
    tc_label   = f'TC-{id_padded}'
    suite      = args.suite or args.scenario or 'TS-001'
    suite_pad  = suite if suite.startswith('TS-') else f'TS-{suite.zfill(3)}'
    notes_lines = [
        f'Suite       : {suite_pad}',
        f'Test Case   : {tc_label}',
        f'Browser     : Chromium (Headed)',
    ]
    if args.notes:
        notes_lines.append(f'Remarks     : {args.notes}')
    return {
        'Suite':          suite_pad,
        'Test Case':      tc_label,
        'Test Name':      f'{suite_pad}/{tc_label}',
        'Target URL':     TARGET_URL,
        'Status':         args.status,
        'Execution Date': date.today().isoformat(),
        'Notes':          '\n'.join(notes_lines),
    }

args   = parse_args()
record = build_record(args)
path   = BASE_DIR / 'E2E_Test_Execution_Summary.xlsx'

if path.exists():
    existing = load_existing_records(path)
else:
    existing = []

existing.append(record)
rebuild_workbook(path, existing)

action = 'Updated' if len(existing) > 1 else 'Created'
print(f'{action} styled report \u2192 {path}  ({len(existing)} row(s))')
